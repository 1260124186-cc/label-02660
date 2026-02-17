"""
Excel导出模块
支持追加和覆盖两种写入模式
"""
import logging
from pathlib import Path
from typing import List, Optional
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

logger = logging.getLogger(__name__)

# Excel列定义
COLUMNS = {
    "source_file": "源文件",
    "award_date": "获奖时间",
    "award_name": "获奖名称",
    "award_org": "授奖单位",
    "award_level": "获奖等级",
    "teachers": "指导老师",
    "students": "获奖学生",
    "confidence": "识别置信度",
}

COLUMN_WIDTHS = {
    "源文件": 30,
    "获奖时间": 14,
    "获奖名称": 35,
    "授奖单位": 30,
    "获奖等级": 12,
    "指导老师": 20,
    "获奖学生": 25,
    "识别置信度": 12,
}

# 样式定义
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
CELL_FONT = Font(name="微软雅黑", size=10)
CELL_ALIGNMENT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)
EVEN_ROW_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")


class ExcelExporter:
    """Excel导出器"""

    def __init__(self):
        pass

    def export(self, data: List[dict], output_path: str, mode: str = "overwrite") -> bool:
        """
        导出数据到Excel
        Args:
            data: 证书信息字典列表
            output_path: 输出文件路径
            mode: 写入模式 "overwrite"(覆盖) 或 "append"(追加)
        Returns:
            是否导出成功
        """
        if not data:
            logger.warning("没有数据可导出")
            return False

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        try:
            df = self._prepare_dataframe(data)

            if mode == "append" and output_path.exists():
                return self._append_to_excel(df, output_path)
            else:
                return self._write_new_excel(df, output_path)

        except Exception as e:
            logger.error(f"Excel导出失败: {e}")
            return False

    def _prepare_dataframe(self, data: List[dict]) -> pd.DataFrame:
        """准备DataFrame"""
        rows = []
        for item in data:
            row = {}
            for key, cn_name in COLUMNS.items():
                row[cn_name] = item.get(key, "")
            rows.append(row)

        df = pd.DataFrame(rows, columns=list(COLUMNS.values()))
        return df

    def _write_new_excel(self, df: pd.DataFrame, output_path: Path) -> bool:
        """写入新Excel文件（带格式）"""
        try:
            df.to_excel(str(output_path), index=False, engine="openpyxl")
            self._apply_styles(output_path, len(df))
            logger.info(f"Excel文件已创建: {output_path}，共 {len(df)} 条记录")
            return True
        except Exception as e:
            logger.error(f"写入Excel失败: {e}")
            return False

    def _append_to_excel(self, df: pd.DataFrame, output_path: Path) -> bool:
        """追加数据到已有Excel文件"""
        try:
            existing_df = pd.read_excel(str(output_path), engine="openpyxl")
            combined_df = pd.concat([existing_df, df], ignore_index=True)
            combined_df.to_excel(str(output_path), index=False, engine="openpyxl")
            self._apply_styles(output_path, len(combined_df))
            logger.info(f"Excel文件已追加: {output_path}，新增 {len(df)} 条，"
                        f"共 {len(combined_df)} 条记录")
            return True
        except Exception as e:
            logger.error(f"追加Excel失败: {e}")
            return False

    def _apply_styles(self, output_path: Path, row_count: int):
        """应用Excel样式"""
        try:
            wb = load_workbook(str(output_path))
            ws = wb.active
            ws.title = "证书识别结果"

            # 设置列宽
            for col_idx, col_name in enumerate(COLUMNS.values(), 1):
                col_letter = chr(64 + col_idx) if col_idx <= 26 else f"A{chr(64 + col_idx - 26)}"
                ws.column_dimensions[col_letter].width = COLUMN_WIDTHS.get(col_name, 15)

            # 表头样式
            for cell in ws[1]:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = HEADER_ALIGNMENT
                cell.border = THIN_BORDER

            # 数据行样式
            for row_idx in range(2, row_count + 2):
                for cell in ws[row_idx]:
                    cell.font = CELL_FONT
                    cell.alignment = CELL_ALIGNMENT
                    cell.border = THIN_BORDER
                    if row_idx % 2 == 0:
                        cell.fill = EVEN_ROW_FILL

            # 冻结首行
            ws.freeze_panes = "A2"

            # 设置行高
            ws.row_dimensions[1].height = 30
            for row_idx in range(2, row_count + 2):
                ws.row_dimensions[row_idx].height = 24

            wb.save(str(output_path))
            logger.debug("Excel样式应用完成")
        except Exception as e:
            logger.warning(f"Excel样式应用失败（不影响数据）: {e}")

    @staticmethod
    def get_default_filename() -> str:
        """生成默认输出文件名"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"证书识别结果_{timestamp}.xlsx"
